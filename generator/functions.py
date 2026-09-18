from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
import openpyxl

_SHEET = "DISCIPLINE_FUNCTION"


@dataclass(frozen=True)
class FunctionInfo:
    code: str
    category: str
    order: int


def _find_workbook(pack_dir: Path) -> Path | None:
    for xlsx in sorted(pack_dir.rglob("*.xlsx")):
        try:
            wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
        except Exception:
            continue
        try:
            if _SHEET in wb.sheetnames:
                return xlsx
        finally:
            wb.close()
    return None


def read_discipline_functions(pack_dir, discipline: str) -> list[FunctionInfo]:
    if pack_dir is None:
        return []
    pack_dir = Path(pack_dir)
    wb_path = _find_workbook(pack_dir)
    if wb_path is None:
        return []
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    try:
        ws = wb[_SHEET]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return []
        idx = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
        need = ("Function", "Discipline", "Order", "Category", "Partic")
        if not all(k in idx for k in need):
            return []
        out: list[FunctionInfo] = []
        for raw in rows:
            if raw is None:
                continue
            def cell(name):
                i = idx[name]
                return "" if i >= len(raw) or raw[i] is None else str(raw[i]).strip()
            if cell("Discipline") != discipline or cell("Partic").upper() != "Y":
                continue
            try:
                order = int(float(cell("Order"))) if cell("Order") else 0
            except ValueError:
                order = 0
            out.append(FunctionInfo(code=cell("Function"),
                                    category=cell("Category").upper(),
                                    order=order))
        out.sort(key=lambda f: f.order)
        return out
    finally:
        wb.close()
